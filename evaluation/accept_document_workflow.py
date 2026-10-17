"""Run the real-model correction/export workflow with fully isolated storage.

Usage: C:/pycharm/python.exe -B -X utf8 evaluation/accept_document_workflow.py
No phone is connected or tested. The report stores IDs, hashes and metrics;
only this run's final three exports are retained under output/pdf/workflow-*.
"""
from __future__ import annotations

import copy
from contextlib import closing
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import time
import uuid

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "reports/document_workflow_acceptance.json"
EXPECTED_SHA = "d3b7f29d89e43c24ad812e2f3302d5f55a9b8f2da23db8163116cef929e3cdbe"


def digest(data):
    return hashlib.sha256(data).hexdigest()


def normalize(text):
    return "".join(text.split())


def user_storage_state():
    databases = {}
    for suffix in ("", "-wal", "-shm", "-journal"):
        path = ROOT / ("snapscan.db" + suffix)
        databases[path.name] = digest(path.read_bytes()) if path.is_file() else None
    uploads = ROOT / "uploads"
    files = {}
    if uploads.is_dir():
        for path in uploads.rglob("*"):
            if path.is_file():
                stat = path.stat()
                files[path.relative_to(uploads).as_posix()] = [stat.st_size, stat.st_mtime_ns]
    return {"databases": databases, "uploads": files}


def main():
    started = time.perf_counter()
    before = user_storage_state()
    previous_model = os.environ.get("SNAPSCAN_OCR_MODEL")
    report = {
        "created_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "scope": "Real rec_head_v1 through FastAPI TestClient; not a mobile-device test",
        "real_device_verified": False, "test_set_read": False,
        "private_full_text_in_report": False, "checks": {}, "artifacts": [],
    }

    def check(name, condition, **metrics):
        ok = bool(condition)
        report["checks"][name] = {"pass": ok, **metrics}
        print(f"{'PASS' if ok else 'FAIL'} {name}", flush=True)
        if not ok:
            raise AssertionError(name)

    try:
        check("fresh_process_before_database_import", "app.main" not in sys.modules and "app.core.db" not in sys.modules)
        with tempfile.TemporaryDirectory(prefix="snapscan-document-workflow-") as directory:
            temporary = Path(directory)
            uploads = temporary / "uploads"
            uploads.mkdir()
            sys.path.insert(0, str(ROOT / "backend"))
            from app.core import config
            config.DB_PATH = str(temporary / "workflow.sqlite3")
            config.UPLOAD_DIR = str(uploads)
            os.environ["SNAPSCAN_OCR_MODEL"] = "rec_head_v1"
            from app.main import app
            from app.core import db
            from app.routers import ocr, export_router
            from app.services import ocr_engine
            from fastapi.testclient import TestClient
            import cv2
            import numpy as np
            from PIL import Image
            from pypdf import PdfReader
            from pypdf.generic import ContentStream
            from openpyxl import load_workbook

            check("isolated_database_and_all_output_routes",
                  Path(db.DB_PATH).parent == temporary and Path(ocr.UPLOAD_DIR) == uploads
                  and Path(export_router.UPLOAD_DIR) == uploads
                  and any(getattr(route, "path", None) == "/uploads"
                          and Path(route.app.directory) == uploads for route in app.routes))
            sources = [json.loads(line) for line in (ROOT / "training/data/val.jsonl").read_text(encoding="utf-8").splitlines()]
            source = next(row for row in sources if row["source_image_id"] == "real-071.jpg")
            image_path = (ROOT / "data/real" / source["source_image_path"]).resolve()
            check("fixture_is_complete_dev_image", source.get("source_split") == "dev"
                  and source.get("split") == "val"
                  and image_path.is_relative_to((ROOT / "data/real/dev").resolve()))
            raw = image_path.read_bytes()
            check("fixture_matches_validation_source_hash", digest(raw) == source["source_sha256"])
            decoded = cv2.imdecode(np.frombuffer(raw, np.uint8), cv2.IMREAD_COLOR)
            height, width = decoded.shape[:2]
            marker = "已校对验收" + uuid.uuid4().hex[:10]
            report["sample"] = {"image_id": "real-071.jpg", "source_split": "dev",
                                "width": int(width), "height": int(height), "sha256": digest(raw)}
            report["correction_marker"] = marker
            final_exports = []
            ocr_engine._singleton = None

            with TestClient(app) as client:
                response = client.get("/api/health")
                health = response.json()
                check("health_identifies_real_trained_model", response.status_code == 200
                      and health.get("ocr", {}).get("ready") is True
                      and health["ocr"]["model"] == {"id": "rec_head_v1", "sha256": EXPECTED_SHA})
                engine = ocr_engine.get_engine()
                session_path = Path(engine.engine.text_rec.session.session._model_path).resolve()
                check("actual_ort_weights_match_reported_model",
                      session_path == (ROOT / "backend/models/rec_head_v1.onnx").resolve()
                      and digest(session_path.read_bytes()) == EXPECTED_SHA,
                      sha256=EXPECTED_SHA)
                response = client.post("/api/ocr/recognize", files={"file": ("real-071.jpg", raw, "image/jpeg")})
                original = response.json()
                check("real_ocr_post", response.status_code == 200 and original.get("ok") is True
                      and original.get("model") == {"id": "rec_head_v1", "sha256": EXPECTED_SHA}
                      and len(original.get("boxes", [])) > 0 and original.get("revision") == 0,
                      boxes=len(original.get("boxes", [])), text_chars=len(original.get("text", "")),
                      server_ms=original.get("elapsed_ms"))
                sid = original["id"]
                response = client.get(original["image"])
                png = response.content
                stored = cv2.imdecode(np.frombuffer(png, np.uint8), cv2.IMREAD_COLOR)
                check("saved_png_download_matches_decoded_ocr_input",
                      response.status_code == 200 and png.startswith(b"\x89PNG\r\n\x1a\n")
                      and np.array_equal(stored, decoded)
                      and (original["image_width"], original["image_height"]) == (width, height),
                      png_bytes=len(png), png_sha256=digest(png))
                response = client.get(f"/api/history/{sid}")
                initial = response.json()
                document = initial.get("document", {})
                check("history_has_normalized_document_and_legacy_item", response.status_code == 200
                      and initial.get("ok") is True and isinstance(initial["item"]["boxes"], str)
                      and document.get("boxes") == original["boxes"]
                      and document.get("original_boxes") == original["boxes"]
                      and document.get("model") == original["model"]
                      and document.get("image") == original["image"])

                texts = [box["text"] for box in original["boxes"]]
                candidates = [index for index, text in enumerate(texts)
                              if len(normalize(text)) >= 4
                              and normalize(text) not in normalize("\n".join(texts[:index] + texts[index+1:]))
                              and normalize(text) not in marker]
                check("unique_line_available_for_replacement", bool(candidates) and marker not in original["text"])
                index = candidates[0]
                old_line = texts[index]
                texts[index] = marker
                expected_text = "\n".join(texts)
                expected_boxes = copy.deepcopy(original["boxes"])
                expected_boxes[index]["text"] = marker
                response = client.patch(f"/api/history/{sid}", json={"texts": texts, "revision": 0})
                corrected = response.json().get("document", {})
                check("patch_exactly_one_line_with_revision_increment", response.status_code == 200
                      and corrected.get("revision") == 1 and corrected.get("text") == expected_text
                      and corrected.get("boxes") == expected_boxes,
                      line_index=index, replaced_line_chars=len(old_line), replaced_line_sha256=digest(old_line.encode()),
                      text_chars=len(expected_text))
                check("patch_preserves_original_boxes_image_geometry_and_model",
                      corrected.get("original_boxes") == original["boxes"]
                      and corrected.get("image") == original["image"]
                      and corrected.get("model") == original["model"]
                      and (corrected.get("image_width"), corrected.get("image_height")) == (width, height))
                response = client.get("/api/history/search", params={"q": marker})
                search = response.json()
                check("corrected_word_is_searchable", response.status_code == 200 and search.get("ok") is True
                      and {item["id"] for item in search.get("items", [])} == {sid},
                      method=search.get("method"), hits=len(search.get("items", [])))

                def export(endpoint, mode=None):
                    payload = {"scan_id": sid, "revision": 1, "text": "stale-client-value", "boxes": []}
                    if mode:
                        payload["mode"] = mode
                    response = client.post("/api/export/" + endpoint, json=payload)
                    body = response.json()
                    tag = mode or endpoint
                    check("export_" + tag, response.status_code == 200 and body.get("ok") is True
                          and body.get("revision") == 1 and (mode is None or body.get("mode") == mode))
                    download = client.get(body["file"])
                    check("download_" + tag, download.status_code == 200 and bool(download.content),
                          bytes=len(download.content))
                    return download.content

                searchable = export("pdf", "searchable")
                reader = PdfReader(io.BytesIO(searchable))
                check("searchable_pdf_has_one_page", len(reader.pages) == 1)
                page = reader.pages[0]
                extracted = page.extract_text() or ""
                check("searchable_pdf_contains_only_latest_text", normalize(extracted) == normalize(expected_text)
                      and marker in extracted and normalize(old_line) not in normalize(extracted)
                      and "stale-client-value" not in extracted, extracted_chars=len(extracted))
                check("searchable_pdf_retains_exact_scan_pixels", len(page.images) == 1
                      and page.images[0].image.convert("RGB").tobytes()
                      == Image.fromarray(cv2.cvtColor(decoded, cv2.COLOR_BGR2RGB)).tobytes(),
                      embedded_images=len(page.images))
                check("searchable_pdf_preserves_page_ratio",
                      abs(float(page.mediabox.width) / float(page.mediabox.height) - width/height) < 1e-6,
                      width_pt=float(page.mediabox.width), height_pt=float(page.mediabox.height))
                operations = ContentStream(page.get_contents(), reader).operations
                modes = [int(arguments[0]) for arguments, operator in operations if operator == b"Tr"]
                check("searchable_text_is_invisible", bool(modes) and set(modes) == {3}, text_runs=len(modes))
                final_exports.append(("workflow-real-071-searchable.pdf", searchable, "image-backed searchable PDF"))

                excel = export("excel")
                workbook = load_workbook(io.BytesIO(excel), read_only=True, data_only=False)
                sheet = workbook.active
                rows = list(sheet.iter_rows(values_only=True))
                check("excel_rows_are_exactly_latest_ocr_lines", len(rows) == len(expected_boxes) + 1
                      and [row[1] for row in rows[1:]] == texts
                      and all(abs(float(row[2])-box["score"]) < 1e-12
                              and row[3] == str(box["box"]) for row, box in zip(rows[1:], expected_boxes)),
                      rows=len(rows), expected_rows=len(expected_boxes)+1)
                check("excel_replaced_line_is_literal_latest_value", sheet.cell(index+2, 2).value == marker
                      and sheet.cell(index+2, 2).data_type == "s")
                workbook.close()
                final_exports.append(("workflow-real-071-ocr-rows.xlsx", excel, "corrected OCR rows, not table reconstruction"))

                text_pdf = export("pdf", "text")
                reader = PdfReader(io.BytesIO(text_pdf))
                extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
                check("plain_text_pdf_contains_only_latest_text", normalize(extracted) == normalize(expected_text)
                      and marker in extracted and normalize(old_line) not in normalize(extracted)
                      and all(len(page.images) == 0 for page in reader.pages), pages=len(reader.pages))
                final_exports.append(("workflow-real-071-text.pdf", text_pdf, "reflowed text PDF"))

                files_before_conflicts = {path.name for path in uploads.iterdir()}
                for endpoint, mode in (("pdf", "searchable"), ("pdf", "text"), ("excel", None)):
                    payload = {"scan_id": sid, "revision": 0}
                    if mode:
                        payload["mode"] = mode
                    response = client.post("/api/export/" + endpoint, json=payload)
                    check("stale_export_rejected_" + (mode or endpoint), response.status_code == 409)
                response = client.patch(f"/api/history/{sid}", json={"texts": texts, "revision": 0})
                check("stale_correction_rejected", response.status_code == 409)
                check("conflicts_create_no_files", {path.name for path in uploads.iterdir()} == files_before_conflicts)
                response = client.get(f"/api/history/{sid}")
                final = response.json().get("document", {})
                check("history_preserves_original_and_final_revision", response.status_code == 200 and final == corrected
                      and final.get("revision") == 1 and final.get("original_boxes") == original["boxes"])
                check("history_model_is_real_trained_version", final.get("model") == {"id": "rec_head_v1", "sha256": EXPECTED_SHA})
                response = client.get(final["image"])
                check("scan_remains_downloadable_after_all_exports", response.status_code == 200 and digest(response.content) == digest(png))
                with closing(db.get_conn()) as connection:
                    count = connection.execute("SELECT COUNT(*) FROM scans").fetchone()[0]
                check("only_one_isolated_history_record", count == 1, rows=count)

                output = ROOT / "output/pdf"
                output.mkdir(parents=True, exist_ok=True)
                for name, data, kind in final_exports:
                    target = output / name
                    target.write_bytes(data)
                    check("retained_" + name, digest(target.read_bytes()) == digest(data), bytes=len(data))
                    report["artifacts"].append({"path": target.relative_to(ROOT).as_posix(), "kind": kind,
                                                 "bytes": len(data), "sha256": digest(data)})
                check("retained_exactly_three_current_exports", len(final_exports) == len(report["artifacts"]) == 3)
        check("temporary_storage_removed", not temporary.exists())
    except Exception as exc:
        # Do not print/save exception text or raw API bodies; OCR can be private.
        report["checks"]["workflow_exception"] = {"pass": False, "exception_type": type(exc).__name__}
        print(f"FAIL workflow_exception ({type(exc).__name__})", flush=True)
    finally:
        if previous_model is None:
            os.environ.pop("SNAPSCAN_OCR_MODEL", None)
        else:
            os.environ["SNAPSCAN_OCR_MODEL"] = previous_model
        if "app.services.ocr_engine" in sys.modules:
            sys.modules["app.services.ocr_engine"]._singleton = None
        try:
            check("user_database_and_uploads_unchanged", user_storage_state() == before)
        except AssertionError:
            pass
        report["check_count"] = len(report["checks"])
        report["passed_count"] = sum(check["pass"] for check in report["checks"].values())
        report["failed_count"] = report["check_count"] - report["passed_count"]
        report["pass_all"] = report["failed_count"] == 0
        report["duration_seconds"] = round(time.perf_counter()-started, 3)
        REPORT.parent.mkdir(parents=True, exist_ok=True)
        REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({key: report[key] for key in ("check_count", "passed_count", "failed_count", "pass_all", "duration_seconds")}), flush=True)
        print(f"Report: {REPORT}", flush=True)
    return 0 if report["pass_all"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
