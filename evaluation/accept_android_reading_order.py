"""Independently verify captured Android row-layout exports, without rerunning OCR."""
from __future__ import annotations

import hashlib
from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess

from openpyxl import load_workbook
from PIL import Image, ImageChops
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "output/scan/android-device"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def nonspace(text):
    return "".join(text.split())


def main():
    proof_path = FOLDER / "reading-order-proof.json"
    proof = json.loads(proof_path.read_text(encoding="utf-8"))
    artifacts = {}
    for role, entry in proof["artifacts"].items():
        path = FOLDER / Path(entry["path"]).name
        assert path.stat().st_size == entry["bytes"]
        assert digest(path) == entry["sha256"], role
        artifacts[role] = path
    assert len(artifacts) == 7
    document = json.loads(artifacts["layout_json"].read_text(encoding="utf-8"))
    rows = json.loads(artifacts["rows_json"].read_text(encoding="utf-8"))
    text = artifacts["text"].read_text(encoding="utf-8")
    assert document["text"] == text and document["text_rows"] == rows == proof["text_rows"]
    assert document["text_layout"] == "spatial"
    boxes = document["boxes"]
    flat = [index for row in rows for index in row]
    assert sorted(flat) == list(range(len(boxes)))
    expected = "\n".join("  ".join(boxes[index]["text"] for index in row if boxes[index]["text"])
                         for row in rows if any(boxes[index]["text"] for index in row))
    assert text == expected
    assert any(87 in row and 88 in row for row in rows)
    for box, original in zip(boxes, document["original_boxes"]):
        assert box["box"] == original["box"] and box["score"] == original["score"]

    text_pdf = PdfReader(artifacts["text_pdf"])
    assert len(text_pdf.pages) == 1
    extracted = text_pdf.pages[0].extract_text()
    assert extracted.rstrip("\n") == text
    assert len(extracted.splitlines()) == len(rows) == 40
    searchable = PdfReader(artifacts["searchable_pdf"])
    assert len(searchable.pages) == 1
    page = searchable.pages[0]
    assert nonspace(page.extract_text()) == nonspace("".join(boxes[index]["text"] for index in flat))
    assert len(page.images) == 1
    source = Image.open(artifacts["source_image"]).convert("RGB")
    embedded = Image.open(BytesIO(page.images[0].data)).convert("RGB")
    assert source.size == embedded.size
    assert ImageChops.difference(source, embedded).getbbox() is None
    sheet = load_workbook(artifacts["excel"]).active
    assert [sheet.cell(row, 2).value for row in range(2, sheet.max_row + 1)] == [box["text"] for box in boxes]
    assert all(cell.data_type != "f" for row in sheet for cell in row)

    device_source = json.loads((FOLDER / "source-snapshot.json").read_text(encoding="utf-8"))
    release_source = json.loads((ROOT / "build/android-local-source-snapshot.json").read_text(encoding="utf-8"))
    assert device_source == release_source
    computed = hashlib.sha256(json.dumps(device_source["files"], sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    assert computed == device_source["sha256"]
    for name, sha in device_source["files"].items():
        assert digest(ROOT / "mobile/android-local" / name) == sha, name
    stable = ROOT / "releases/SnapScan-offline-local.apk"
    build = json.loads((ROOT / "reports/android_local_build_verification.json").read_text(encoding="utf-8"))
    device = json.loads((ROOT / "reports/android_local_device_verification.json").read_text(encoding="utf-8"))
    assert build["version_code"] == device["version_code"]
    assert build["source_sha256"] == computed == device["source_sha256"]
    versioned = ROOT / build["versioned_apk"]
    assert digest(stable) == digest(versioned)

    renderer = shutil.which("pdftoppm")
    if not renderer:
        raise RuntimeError("pdftoppm is required for visual inspection")
    rendered = []
    for key in ("text_pdf", "searchable_pdf"):
        path = artifacts[key]
        target = path.with_name(path.stem + "-render")
        subprocess.run([renderer, "-f", "1", "-singlefile", "-r", "120", "-png", str(path), str(target)], check=True, capture_output=True)
        rendered.append(target.with_suffix(".png").relative_to(ROOT).as_posix())
    report = {
        "passed": True,
        "source_kind": "prior actual Android OCR boxes; not coordinates recovered from the supplied user PDF",
        "runtime": "Android 15 API35 emulator, captured native DocumentStore and PDFBox outputs",
        "physical_iqoo_phone_tested": False,
        "version_code": build["version_code"],
        "apk_sha256": digest(versioned),
        "source_sha256": computed,
        "proof": proof_path.relative_to(ROOT).as_posix(),
        "proof_sha256": digest(proof_path),
        "captured_artifacts_hashes_match": True,
        "original_boxes": len(boxes),
        "text_rows": len(rows),
        "all_original_indices_once": True,
        "coordinates_and_scores_preserved": True,
        "text_pdf_pages": len(text_pdf.pages),
        "text_pdf_extracted_rows_exactly_match_saved_summary": True,
        "searchable_pdf_pages": len(searchable.pages),
        "searchable_pdf_image_pixels_unchanged": True,
        "searchable_pdf_content_stream_reading_order_verified": True,
        "excel_retains_original_box_order": True,
        "rendered_pages": rendered,
        "visual_review_completed": False,
        "limitations": [
            "The heading was deliberately edited in the save/reopen regression test.",
            "This is geometric row grouping, not table structure or semantic multi-column recovery.",
            "Original OCR mistakes and noise are preserved; this does not measure OCR accuracy.",
            "Readers may insert different whitespace when copying the positioned searchable PDF."
        ]
    }
    (ROOT / "reports/android_reading_order_acceptance.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
