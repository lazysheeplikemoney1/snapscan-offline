"""Independently inspect Android-generated files and render PDFs for visual QA."""
from __future__ import annotations

import hashlib
from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess

from PIL import Image, ImageChops
from pypdf import PdfReader
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output/scan/android-device"


def main():
    document = json.loads((OUTPUT / "receipt-document.json").read_text(encoding="utf-8"))
    original = Image.open(OUTPUT / "receipt-clean.png").convert("RGB")
    persisted = Image.open(OUTPUT / "receipt-ocr-persisted.png").convert("RGB")
    assert original.size == (document["image_width"], document["image_height"]) == (414, 885)
    assert ImageChops.difference(original, persisted).getbbox() is None
    pdf_path = OUTPUT / "receipt-searchable.pdf"
    receipt = PdfReader(pdf_path)
    assert len(receipt.pages) == 1
    page = receipt.pages[0]
    text = page.extract_text()
    assert "离线校对测试" in text and "933" in text
    assert len(page.images) == 1
    embedded = Image.open(BytesIO(page.images[0].data)).convert("RGB")
    assert embedded.size == original.size and ImageChops.difference(original, embedded).getbbox() is None, "The PDF changed the actual scan pixels"
    workbook = load_workbook(OUTPUT / "receipt-ocr-details.xlsx")
    sheet = workbook.active
    rows = [sheet.cell(i, 2).value for i in range(2, sheet.max_row + 1)]
    assert rows == [box["text"] for box in document["boxes"]]
    assert all(cell.data_type != "f" for row in sheet for cell in row), "OCR text interpreted as an Excel formula"
    proof = PdfReader(OUTPUT / "unicode-proof-text.pdf")
    proof_text = "\n".join(page.extract_text() or "" for page in proof.pages)
    for token in ("中文", "❤", "❋", "ɔ", "ʌ", "𣇉", "🙂", "\U0010fffd", "部分字符无内置字形"):
        assert token in proof_text, "Unicode proof lost " + repr(token)
    renderer = shutil.which("pdftoppm")
    if not renderer:
        raise RuntimeError("pdftoppm is needed for actual PDF rendering")
    for name in ("receipt-searchable", "unicode-proof-text"):
        subprocess.run([renderer, "-f", "1", "-singlefile", "-r", "120", "-png", str(OUTPUT / (name + ".pdf")), str(OUTPUT / (name + "-render"))], check=True, capture_output=True)
    (OUTPUT / "receipt-searchable-text.txt").write_text(text, encoding="utf-8")
    (OUTPUT / "unicode-proof-text.txt").write_text(proof_text, encoding="utf-8")
    report = {"passed": True, "source": "actual Android15 emulator-generated files", "physical_phone_tested": False,
              "receipt_size": list(original.size), "saved_ocr_image_pixels_identical": True, "pdf_image_pixels_identical": True,
              "pdf_searchable_chinese_verified": True, "excel_rows_match_saved_revision": True, "excel_rows": len(rows),
              "unicode_roundtrip_verified_independently_with_pypdf": True,
              "renderer": renderer, "rendered_pages": ["output/scan/android-device/receipt-searchable-render.png", "output/scan/android-device/unicode-proof-text-render.png"],
              "visual_review_completed": False,
              "input_sha256": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in OUTPUT.glob("*") if p.suffix in (".pdf", ".xlsx", ".png")}}
    path = ROOT / "reports/android_local_export_acceptance.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: value for key, value in report.items() if key != "input_sha256"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
