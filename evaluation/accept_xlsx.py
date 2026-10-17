# -*- coding: utf-8 -*-
"""XLSX验收：OCR结果导出XLSX（非表格恢复）。输出 xlsx_acceptance.json。"""
import json, os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
import sys
sys.path.insert(0, os.path.join(ROOT, "backend"))
from app.services.export import export_boxes_to_excel
import openpyxl
OUT_X = os.path.join(ROOT, "uploads", "accept_xlsx_check.xlsx")
OUT_JSON = os.path.join(ROOT, "evaluation", "xlsx_acceptance.json")
boxes = [{"text": "卷积神经网络", "score": 0.99, "box": [[0,0],[1,0],[1,1],[0,1]]},
         {"text": "发票金额128.50元", "score": 0.95, "box": [[0,0],[1,0],[1,1],[0,1]]}]
export_boxes_to_excel(boxes, OUT_X)
wb = openpyxl.load_workbook(OUT_X)
ws = wb.active
vals = [c.value for r in ws.iter_rows(min_row=2) for c in r]
res = {"file": OUT_X, "sheet": ws.title, "rows": ws.max_row,
       "has_cn": any("卷积" in str(v) for v in vals),
       "pass": ws.max_row >= 3 and any("卷积" in str(v) for v in vals),
       "scope": "OCR结果导出XLSX，不含表格结构恢复"}
json.dump(res, open(OUT_JSON, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
print(json.dumps(res, ensure_ascii=False), "saved:", OUT_JSON)
