from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from app.main import app
from app.routers import export_router


@pytest.fixture
def exports(monkeypatch, tmp_path):
    image_name = 'scan_' + 'a'*32 + '.png'
    Image.new('RGB', (200, 100), '#d9f0e8').save(tmp_path / image_name)
    document = {'id': 9, 'revision': 2, 'image': '/uploads/'+image_name,
                'image_width': 200, 'image_height': 100,
                'text': '已校对金额 00128.50',
                'boxes': [{'text': '已校对金额 00128.50', 'box': [[5, 5], [195, 5], [195, 40], [5, 40]], 'score': .9}]}
    monkeypatch.setattr(export_router, 'UPLOAD_DIR', str(tmp_path))
    monkeypatch.setattr(export_router, 'load_document', lambda sid: document if sid == 9 else None, raising=False)
    return TestClient(app), document, tmp_path


def output_path(response, root):
    assert response.status_code == 200, response.text
    return root / Path(response.json()['file']).name


def test_document_pdf_uses_saved_scan_and_latest_corrections(exports):
    client, document, root = exports
    response = client.post('/api/export/pdf', json={'scan_id': 9, 'revision': 2, 'mode': 'searchable', 'text': 'stale', 'boxes': []})
    page = PdfReader(output_path(response, root)).pages[0]
    assert len(page.images) == 1
    assert document['text'] in page.extract_text()
    assert 'stale' not in page.extract_text()
    assert response.json()['mode'] == 'searchable'


def test_excel_uses_same_saved_revision(exports):
    client, document, root = exports
    response = client.post('/api/export/excel', json={'scan_id': 9, 'revision': 2, 'text': 'stale'})
    assert load_workbook(output_path(response, root)).active['B2'].value == document['text']


@pytest.mark.parametrize('endpoint', ['pdf', 'excel'])
def test_stale_export_is_rejected_without_creating_a_file(exports, endpoint):
    client, document, root = exports
    response = client.post('/api/export/'+endpoint, json={'scan_id': 9, 'revision': 1})
    assert response.status_code == 409
    assert not list(root.glob('export_*'))


def test_unknown_scan_is_not_silently_exported_as_blank(exports):
    client, document, root = exports
    assert client.post('/api/export/pdf', json={'scan_id': 99, 'revision': 0}).status_code == 404


def test_searchable_pdf_requires_a_saved_image(exports):
    client, document, root = exports
    document['image'] = None
    response = client.post('/api/export/pdf', json={'scan_id': 9, 'revision': 2, 'mode': 'searchable'})
    assert response.status_code == 409
    assert '扫描图' in response.json()['detail']
    response = client.post('/api/export/pdf', json={'scan_id': 9, 'revision': 2, 'mode': 'text'})
    assert document['text'] in PdfReader(output_path(response, root)).pages[0].extract_text()


@pytest.mark.parametrize('path', ['/uploads/../private.png', 'https://example.com/page.png', '/uploads/export_other.png'])
def test_image_reference_cannot_choose_an_arbitrary_file(exports, path):
    client, document, root = exports
    document['image'] = path
    assert client.post('/api/export/pdf', json={'scan_id': 9, 'revision': 2}).status_code == 409
    assert not list(root.glob('export_*'))


def test_legacy_text_only_excel_contains_the_requested_text(exports):
    client, document, root = exports
    response = client.post('/api/export/excel', json={'text': '=1+2\n00128.50'})
    sheet = load_workbook(output_path(response, root)).active
    assert sheet['B2'].value == '=1+2' and sheet['B2'].data_type == 's'
    assert sheet['B3'].value == '00128.50'


def test_searchable_export_does_not_accept_client_supplied_image(exports):
    client, document, root = exports
    assert client.post('/api/export/pdf', json={'mode': 'searchable', 'image': '/private.png'}).status_code == 422
