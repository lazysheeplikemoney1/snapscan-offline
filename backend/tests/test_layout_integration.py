"""Layout stays derived from immutable box indices across persistence/exports."""
import copy
import json
from pathlib import Path

import cv2
import numpy as np
import pytest
from fastapi import FastAPI
from fastapi.staticfiles import StaticFiles
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from app.core import config, db
from app.routers import ocr, history, export_router
from app.services import documents, export
from app.services.ocr_engine import RapidEngine


BOXES = [
    {"text": "12.00", "score": .91, "box": [[100, 15], [140, 15], [140, 25], [100, 25]]},
    {"text": "下一行", "score": .82, "box": [[5, 55], [80, 55], [80, 65], [5, 65]]},
    {"text": "项目", "score": .93, "box": [[5, 10], [80, 10], [80, 30], [5, 30]]},
]
EXPECTED = "项目  12.00\n下一行"


@pytest.fixture
def runtime(monkeypatch, tmp_path):
    uploads = tmp_path / "uploads"
    uploads.mkdir()
    monkeypatch.setattr(db, "DB_PATH", str(tmp_path / "layout.sqlite3"))
    db.init_db()
    monkeypatch.setattr(ocr, "UPLOAD_DIR", str(uploads))
    monkeypatch.setattr(export_router, "UPLOAD_DIR", str(uploads))
    engine = RapidEngine.__new__(RapidEngine)
    engine.model_info = {"id": "layout-fixture", "sha256": "a"*64}
    engine.engine = lambda image: ([(item["box"], item["text"], item["score"]) for item in copy.deepcopy(BOXES)], None)
    monkeypatch.setattr(ocr, "get_engine", lambda: engine)
    app = FastAPI()
    app.mount("/uploads", StaticFiles(directory=uploads), name="uploads")
    for router in (ocr.router, history.router, export_router.router):
        app.include_router(router)
    with TestClient(app) as client:
        yield client, uploads


def recognize(client):
    image = np.full((100, 160, 3), 235, np.uint8)
    response = client.post("/api/ocr/recognize", files={"file": ("layout.png", cv2.imencode(".png", image)[1].tobytes(), "image/png")})
    assert response.status_code == 200, response.text
    return response.json()


def raw_row(sid):
    connection = db.get_conn()
    try:
        return dict(connection.execute("SELECT * FROM scans WHERE id=?", (sid,)).fetchone())
    finally:
        connection.close()


def test_real_engine_adapter_returns_row_indices_without_reordering_boxes(runtime):
    client, uploads = runtime
    body = recognize(client)
    assert body["text"] == EXPECTED
    assert body["text_rows"] == [[2, 0], [1]] and body["text_layout"] == "spatial"
    assert body["boxes"] == BOXES
    stored = raw_row(body["id"])
    assert stored["text"] == EXPECTED
    assert json.loads(stored["boxes"]) == BOXES == json.loads(stored["original_boxes"])


def test_old_history_read_derives_rows_without_writing_db_or_revision(runtime):
    client, uploads = runtime
    connection = db.get_conn()
    sid = connection.execute("INSERT INTO scans(text,boxes,revision) VALUES(?,?,7)", ("legacy text", json.dumps(BOXES))).lastrowid
    connection.commit()
    connection.close()
    before = raw_row(sid)
    response = client.get(f"/api/history/{sid}").json()
    document = response["document"]
    assert document["text"] == EXPECTED and document["text_rows"] == [[2, 0], [1]]
    assert document["revision"] == 7 and document["boxes"] == BOXES
    assert response["item"]["text"] == "legacy text"
    assert raw_row(sid) == before
    assert documents.document_from_row({"text": "无框旧文本", "boxes": "[]"})["text"] == "无框旧文本"


def test_correction_uses_original_indices_and_updates_search_with_spatial_text(runtime):
    client, uploads = runtime
    body = recognize(client)
    texts = ["18.50", "下一行已校对", "项目已改"]
    response = client.patch(f"/api/history/{body['id']}", json={"texts": texts, "revision": 0})
    assert response.status_code == 200, response.text
    document = response.json()["document"]
    assert document["text"] == "项目已改  18.50\n下一行已校对" and document["revision"] == 1
    assert document["text_rows"] == [[2, 0], [1]]
    assert document["boxes"] == [{**item, "text": text} for item, text in zip(BOXES, texts)]
    assert document["original_boxes"] == BOXES
    assert raw_row(body["id"])["text"] == document["text"]
    found = client.get("/api/history/search", params={"q": "项目已改  18.50"}).json()["items"]
    assert body["id"] in {item["id"] for item in found}
    assert client.patch(f"/api/history/{body['id']}", json={"texts": texts, "revision": 0}).status_code == 409


def test_text_pdf_with_boxes_uses_current_box_text_and_reading_rows(runtime):
    client, uploads = runtime
    response = client.post("/api/export/pdf", json={"mode": "text", "text": "stale summary", "boxes": BOXES})
    assert response.status_code == 200, response.text
    text = PdfReader(uploads / Path(response.json()["file"]).name).pages[0].extract_text()
    assert EXPECTED in text and "stale summary" not in text


def test_anonymous_no_box_free_text_is_not_replaced(runtime):
    client, uploads = runtime
    text = "独立草稿\n原样  两个空格"
    response = client.post("/api/export/pdf", json={"mode": "text", "text": text})
    assert response.status_code == 200
    assert text in PdfReader(uploads / Path(response.json()["file"]).name).pages[0].extract_text()


def test_searchable_pdf_content_stream_uses_reading_order_and_keeps_the_image(tmp_path):
    image = tmp_path / "scan.png"
    Image.new("RGB", (160, 100), "#fff5dc").save(image)
    output = tmp_path / "ordered.pdf"
    export.export_searchable_pdf(image, copy.deepcopy(BOXES), output)
    page = PdfReader(output).pages[0]
    text = page.extract_text()
    assert text.index("项目") < text.index("12.00") < text.index("下一行")
    assert len(page.images) == 1


def test_excel_detail_keeps_original_box_order_and_indices(runtime):
    client, uploads = runtime
    response = client.post("/api/export/excel", json={"text": EXPECTED, "boxes": BOXES})
    assert response.status_code == 200, response.text
    sheet = load_workbook(uploads / Path(response.json()["file"]).name).active
    assert [sheet.cell(i, 1).value for i in range(2, 5)] == [1, 2, 3]
    assert [sheet.cell(i, 2).value for i in range(2, 5)] == [item["text"] for item in BOXES]
